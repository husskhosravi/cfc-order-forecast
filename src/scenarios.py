"""Step 3 - Scenario ("what-if") analysis.

Scenarios follow the brief: increased promotional activity and an unexpected
market shift. All scenarios start from the FY25 baseline in outputs/forecast_daily.csv.

Run:  python src/scenarios.py
Writes: outputs/scenario_summary.csv, outputs/scenario_monthly.csv,
        outputs/scenario_tool.xlsx (formula-driven, editable), charts/scenario_*.png
"""
import json

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import BLUE, CHARTS, GREY, OUT, RED

Z80 = 1.2816
CAPACITY = 550  # illustrative daily planning capacity used to count pressure days

SCENARIOS = {
    "Base case": {
        "promo_uplift": 0.25, "extra_promos": 0, "market_shift": 0.0, "shift_start": "2025-01-01",
        "description": "Assumed promotional calendar (17 days) at +25%, no market shift.",
    },
    "Scenario 1: Increased promotions": {
        "promo_uplift": 0.50, "extra_promos": 1, "market_shift": 0.0, "shift_start": "2025-01-01",
        "description": "Stronger promotions (+50%) plus an extra 2-day promotion each month.",
    },
    "Scenario 2: Market downturn": {
        "promo_uplift": 0.25, "extra_promos": 0, "market_shift": -0.10, "shift_start": "2025-01-01",
        "description": "Unexpected 10% drop in underlying demand from 1 Jan 2025 (e.g. a competitor "
                       "price war or cost-of-living pressure); promotions unchanged.",
    },
}


def extra_promo_days(daily: pd.DataFrame) -> pd.Series:
    """First Friday and Saturday of every month that is not already a promotion day."""
    flag = pd.Series(0, index=daily.index)
    for _, month in daily.groupby(daily.index.to_period("M")):
        fri = month[month.index.dayofweek == 4].index[0]
        for d in (fri, fri + pd.Timedelta(days=1)):
            if d in daily.index and daily.loc[d, "is_promo"] == 0:
                flag[d] = 1
    return flag


def apply(daily, p):
    shift = np.where(daily.index >= pd.Timestamp(p["shift_start"]), p["market_shift"], 0.0)
    promo_flag = np.maximum(daily["is_promo"], daily["extra_candidate"] * p["extra_promos"])
    return daily["base_forecast"] * (1 + shift) * (1 + p["promo_uplift"] * promo_flag)


def main() -> None:
    daily = pd.read_csv(OUT / "forecast_daily.csv", index_col="date", parse_dates=True)
    daily["promo_type"] = daily["promo_type"].fillna("")
    model = json.loads((OUT / "model_summary.json").read_text())
    sigma = model["daily_sigma_out_of_sample"]
    base_level = model["fy25"]["baseline_daily"]
    daily["extra_candidate"] = extra_promo_days(daily)

    upper_factor = 1 + Z80 * sigma / base_level
    results, monthly = [], {}
    for name, p in SCENARIOS.items():
        s = apply(daily, p)
        monthly[name] = s.groupby(s.index.to_period("M").astype(str)).sum()
        promo_days = int(np.maximum(daily["is_promo"], daily["extra_candidate"] * p["extra_promos"]).sum())
        results.append({
            "scenario": name,
            "total_orders": round(float(s.sum())),
            "avg_daily": round(float(s.mean()), 1),
            "peak_day_expected": round(float(s.max())),
            "peak_day_80_upper": round(float(s.max() * upper_factor)),
            "promo_days": promo_days,
            "days_80_upper_above_capacity": int((s * upper_factor > CAPACITY).sum()),
        })
    res = pd.DataFrame(results).set_index("scenario")
    base_total = res.loc["Base case", "total_orders"]
    res["vs_base_orders"] = res["total_orders"] - base_total
    res["vs_base_pct"] = ((res["total_orders"] / base_total - 1) * 100).round(1)
    res.to_csv(OUT / "scenario_summary.csv")
    mdf = pd.DataFrame(monthly).round(0)
    mdf.index.name = "month"
    mdf.to_csv(OUT / "scenario_monthly.csv")
    print(res.to_string())

    # Chart: monthly orders by scenario
    fig, ax = plt.subplots(figsize=(11, 4.2))
    labels = pd.to_datetime(mdf.index).strftime("%b %y")
    x = np.arange(len(mdf))
    w = 0.27
    for i, (name, colour) in enumerate(zip(mdf.columns, [BLUE, RED, GREY])):
        ax.bar(x + (i - 1) * w, mdf[name], w, label=name, color=colour)
    ax.set_xticks(x, labels)
    ax.set_ylim(mdf.values.min() * 0.85, mdf.values.max() * 1.05)
    ax.set_ylabel("Orders per month")
    ax.set_title("FY25 monthly orders by scenario")
    ax.legend(frameon=False, ncol=3, loc="upper left")
    fig.savefig(CHARTS / "scenario_monthly.png")
    plt.close(fig)

    build_workbook(daily, sigma, base_level)


# ----------------------------------------------------------------- Excel tool

def build_workbook(daily: pd.DataFrame, sigma: float, base_level: float) -> None:
    wb = Workbook()
    font = "Arial"
    bold = Font(name=font, bold=True)
    normal = Font(name=font)
    input_font = Font(name=font, color="0000FF")
    link_font = Font(name=font, color="008000")
    head_fill = PatternFill("solid", fgColor="2F5D8A")
    head_font = Font(name=font, bold=True, color="FFFFFF")
    input_fill = PatternFill("solid", fgColor="FFFF00")
    thin = Border(bottom=Side(style="thin", color="BFBFBF"))

    def header(ws, row, values):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    # --- Read me
    ws = wb.active
    ws.title = "Read_Me"
    lines = [
        ("FY25 CFC order scenario tool", bold),
        ("", normal),
        ("How to use", bold),
        ("1. Change the yellow cells on the Inputs sheet (blue text = editable inputs).", normal),
        ("2. The Daily sheet recalculates every scenario for each day of FY25.", normal),
        ("3. The Summary sheet shows annual and monthly results and the change versus the base case.", normal),
        ("", normal),
        ("Scenario formula (per day)", bold),
        ("Orders = Baseline x (1 + market shift if date >= shift start) x (1 + promo uplift x promo flag)", normal),
        ("Promo flag = 1 on calendar promotion days, or on extra monthly promotion days when that scenario "
         "includes them.", normal),
        ("", normal),
        ("Sources", bold),
        ("Baseline and prediction error come from src/forecast.py (outputs/forecast_daily.csv).", normal),
        ("Promotion dates and uplift percentages are planning assumptions: FY24 data shows no measurable "
         "promotional uplift.", normal),
    ]
    for r, (text, f) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=text).font = f
    ws.column_dimensions["A"].width = 110

    # --- Inputs
    inp = wb.create_sheet("Inputs")
    names = list(SCENARIOS)
    header(inp, 1, ["Parameter"] + names + ["Notes"])
    params = [
        ("Promotional uplift (%)", "promo_uplift", "0.0%",
         "Applied to promotion days. Assumption: no uplift is visible in FY24 data."),
        ("Include extra monthly promotions (1 = yes, 0 = no)", "extra_promos", "0",
         "Adds the first Friday and Saturday of each month as promotion days."),
        ("Market shift in underlying demand (%)", "market_shift", "0.0%",
         "Applied to the baseline from the shift start date onwards."),
        ("Market shift start date", "shift_start", "dd-mmm-yyyy", ""),
    ]
    for r, (label, key, fmt, note) in enumerate(params, 2):
        inp.cell(row=r, column=1, value=label).font = normal
        for c, name in enumerate(names, 2):
            v = SCENARIOS[name][key]
            if key == "shift_start":
                v = pd.Timestamp(v).to_pydatetime()
            cell = inp.cell(row=r, column=c, value=v)
            cell.font, cell.fill, cell.number_format = input_font, input_fill, fmt
        inp.cell(row=r, column=len(names) + 2, value=note).font = normal
    r0 = len(params) + 3
    inp.cell(row=r0, column=1, value="Model constants (from src/forecast.py)").font = bold
    consts = [
        ("Baseline orders per day", base_level, "0.0", "Holt-Winters level for FY25."),
        ("Daily prediction error (orders, out-of-sample RMSE)", sigma, "0.0",
         "Rolling-origin backtest, 6 x 30 days."),
        ("z-value for 80% interval", Z80, "0.0000", ""),
        ("Planning capacity (orders per day)", CAPACITY, "#,##0",
         "Illustrative. Replace with actual CFC capacity."),
    ]
    for i, (label, v, fmt, note) in enumerate(consts, r0 + 1):
        inp.cell(row=i, column=1, value=label).font = normal
        cell = inp.cell(row=i, column=2, value=v)
        cell.font, cell.number_format = input_font, fmt
        if "capacity" in label.lower():
            cell.fill = input_fill
        inp.cell(row=i, column=len(names) + 2, value=note).font = normal
    base_ref, sigma_ref, z_ref, cap_ref = (f"Inputs!$B${r0 + 1}", f"Inputs!$B${r0 + 2}",
                                           f"Inputs!$B${r0 + 3}", f"Inputs!$B${r0 + 4}")
    inp.column_dimensions["A"].width = 52
    for c in range(2, len(names) + 2):
        inp.column_dimensions[get_column_letter(c)].width = 22
    inp.column_dimensions[get_column_letter(len(names) + 2)].width = 70

    # --- Daily
    d = wb.create_sheet("Daily")
    cols = ["Date", "Month", "Baseline", "Promotion", "Calendar promo (1/0)",
            "Extra monthly promo candidate (1/0)"] + names
    header(d, 1, cols)
    n = len(daily)
    for i, (dt, row) in enumerate(daily.iterrows(), 2):
        d.cell(row=i, column=1, value=dt.to_pydatetime()).number_format = "dd-mmm-yyyy"
        d.cell(row=i, column=2, value=f'=TEXT(A{i},"yyyy-mm")')
        d.cell(row=i, column=3, value=round(float(row["base_forecast"]), 1)).number_format = "0.0"
        d.cell(row=i, column=4, value=row["promo_type"])
        d.cell(row=i, column=5, value=int(row["is_promo"]))
        d.cell(row=i, column=6, value=int(row["extra_candidate"]))
        for j, _ in enumerate(names):
            col = get_column_letter(j + 2)  # column in Inputs
            f = (f"=C{i}*(1+IF(A{i}>=Inputs!${col}$5,Inputs!${col}$4,0))"
                 f"*(1+Inputs!${col}$2*MAX(E{i},F{i}*Inputs!${col}$3))")
            c = d.cell(row=i, column=7 + j, value=f)
            c.number_format = "#,##0.0"
    for c in range(1, len(cols) + 1):
        d.column_dimensions[get_column_letter(c)].width = 16 if c > 1 else 14
    d.column_dimensions["D"].width = 34
    d.freeze_panes = "B2"
    last = n + 1

    # --- Summary
    s = wb.create_sheet("Summary")
    header(s, 1, ["Metric"] + names)
    metrics = [
        ("Total FY25 orders", "=SUM(Daily!{c}2:{c}%d)" % last, "#,##0"),
        ("Change vs base case (orders)", "={c0}2-$B$2", "#,##0;(#,##0);-"),
        ("Change vs base case (%)", "={c0}2/$B$2-1", "0.0%;(0.0%);-"),
        ("Average orders per day", "=AVERAGE(Daily!{c}2:{c}%d)" % last, "#,##0.0"),
        ("Peak day (expected)", "=MAX(Daily!{c}2:{c}%d)" % last, "#,##0"),
        ("Peak day (80% upper bound)", "={c0}6*(1+%s*%s/%s)" % (z_ref, sigma_ref, base_ref), "#,##0"),
        ("Days where 80% upper bound exceeds planning capacity",
         "=COUNTIF(Daily!{c}2:{c}%d,\">\"&(%s/(1+%s*%s/%s)))" % (last, cap_ref, z_ref, sigma_ref, base_ref),
         "#,##0"),
    ]
    for r, (label, tmpl, fmt) in enumerate(metrics, 2):
        s.cell(row=r, column=1, value=label).font = normal
        for j, _ in enumerate(names):
            daily_col = get_column_letter(7 + j)
            sum_col = get_column_letter(2 + j)
            cell = s.cell(row=r, column=2 + j, value=tmpl.format(c=daily_col, c0=sum_col))
            cell.number_format = fmt
            cell.border = thin
    mrow = len(metrics) + 4
    s.cell(row=mrow - 1, column=1, value="Monthly orders").font = bold
    header(s, mrow, ["Month"] + names)
    months = sorted(daily.index.to_period("M").astype(str).unique())
    for i, m in enumerate(months, mrow + 1):
        s.cell(row=i, column=1, value=m)
        for j, _ in enumerate(names):
            col = get_column_letter(7 + j)
            s.cell(row=i, column=2 + j,
                   value=f'=SUMIFS(Daily!{col}$2:{col}${last},Daily!$B$2:$B${last},$A{i})').number_format = "#,##0"
    s.column_dimensions["A"].width = 52
    for c in range(2, len(names) + 2):
        s.column_dimensions[get_column_letter(c)].width = 26
    s.cell(row=2, column=1).comment = Comment("Sum of daily scenario orders on the Daily sheet.", "model")

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.font and cell.font.name != font:
                    cell.font = Font(name=font, bold=cell.font.bold, color=cell.font.color)
    wb.move_sheet("Summary", offset=-2)
    wb.save(OUT / "scenario_tool.xlsx")
    print("Saved outputs/scenario_tool.xlsx")


if __name__ == "__main__":
    main()
